import json
from pathlib import Path

from dubois.report import result_path, write_jsonl, write_txt, write_xlsx
from dubois.result import QueryResult, QueryStatus
from dubois.util import safe_filename


def _sample(username: str = "alice") -> dict:
    claimed = QueryResult(username, "GitHub", "https://github.com/alice", QueryStatus.CLAIMED, query_time=0.2)
    missing = QueryResult(username, "GitLab", "https://gitlab.com/alice", QueryStatus.AVAILABLE, query_time=None)
    return {
        "GitHub": {
            "url_main": "https://github.com/",
            "url_user": "https://github.com/alice",
            "status": claimed,
            "http_status": 200,
            "response_text": "",
        },
        "GitLab": {
            "url_main": "https://gitlab.com/",
            "url_user": "https://gitlab.com/alice",
            "status": missing,
            "http_status": 404,
            "response_text": "",
        },
    }


def test_safe_filename_windows_reserved_and_trailing_dot():
    assert safe_filename("alice.", ".txt") == "alice.txt"
    assert safe_filename("con", ".txt") == "_con.txt"
    assert safe_filename("a/b:c", ".csv") == "a_b_c.csv"


def test_result_path_uses_folder(tmp_path: Path):
    path = result_path("alice.", str(tmp_path), ".xlsx")
    assert path.endswith("alice.xlsx")
    assert str(tmp_path) in path


def test_write_txt_and_jsonl(tmp_path: Path):
    results = _sample()
    txt = tmp_path / "alice.txt"
    write_txt(str(txt), results)
    body = txt.read_text(encoding="utf-8")
    assert "https://github.com/alice" in body
    assert "Total Websites Username Detected On : 1" in body

    jsonl = tmp_path / "alice.jsonl"
    write_jsonl(str(jsonl), "alice", results)
    rows = [json.loads(line) for line in jsonl.read_text(encoding="utf-8").splitlines()]
    assert {row["site"] for row in rows} == {"GitHub", "GitLab"}
    claimed = next(row for row in rows if row["site"] == "GitHub")
    assert claimed["status"] == "Claimed"
    missing = next(row for row in rows if row["site"] == "GitLab")
    assert missing["query_time"] is None


def test_xlsx_none_query_time_and_folder(tmp_path: Path):
    results = _sample()
    path = result_path("alice", str(tmp_path), ".xlsx")
    write_xlsx(path, "alice", results, print_found=True, print_all=True)
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    values = [tuple(cell.value for cell in row) for row in ws.iter_rows(min_row=2)]
    times = {row[1]: row[6] for row in values}
    assert times["GitHub"] == 0.2
    assert times["GitLab"] in (None, "")
