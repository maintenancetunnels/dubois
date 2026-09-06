"""Write search results to txt / csv / xlsx / jsonl. No pandas."""
from __future__ import annotations

import csv
import json
import os
from typing import Any

from dubois.result import QueryStatus
from dubois.util import safe_filename


def _claimed_only(results: dict, print_found: bool, print_all: bool) -> bool:
    return print_found and not print_all


def result_path(username: str, folder: str | None, ext: str, explicit: str | None = None) -> str:
    if explicit:
        return explicit
    name = safe_filename(username, ext)
    if folder:
        os.makedirs(folder, exist_ok=True)
        return os.path.join(folder, name)
    return name


def write_txt(path: str, results: dict[str, dict[str, Any]]) -> int:
    exists_counter = 0
    with open(path, "w", encoding="utf-8") as file:
        for website_name in results:
            dictionary = results[website_name]
            status = dictionary.get("status")
            if status is not None and status.status == QueryStatus.CLAIMED:
                exists_counter += 1
                file.write(dictionary["url_user"] + "\n")
        file.write(f"Total Websites Username Detected On : {exists_counter}\n")
    return exists_counter


def write_csv(
    path: str,
    username: str,
    results: dict[str, dict[str, Any]],
    *,
    print_found: bool = True,
    print_all: bool = False,
) -> None:
    with open(path, "w", newline="", encoding="utf-8") as csv_report:
        writer = csv.writer(csv_report)
        writer.writerow(
            [
                "username",
                "name",
                "url_main",
                "url_user",
                "exists",
                "http_status",
                "response_time_s",
            ]
        )
        for site in results:
            if (
                _claimed_only(results, print_found, print_all)
                and results[site]["status"].status != QueryStatus.CLAIMED
            ):
                continue
            response_time_s = results[site]["status"].query_time
            if response_time_s is None:
                response_time_s = ""
            writer.writerow(
                [
                    username,
                    site,
                    results[site]["url_main"],
                    results[site]["url_user"],
                    str(results[site]["status"].status),
                    results[site]["http_status"],
                    response_time_s,
                ]
            )


def write_xlsx(
    path: str,
    username: str,
    results: dict[str, dict[str, Any]],
    *,
    print_found: bool = True,
    print_all: bool = False,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "sheet1"
    ws.append(
        [
            "username",
            "name",
            "url_main",
            "url_user",
            "exists",
            "http_status",
            "response_time_s",
        ]
    )
    link_font = Font(color="0563C1", underline="single")
    for site in results:
        if (
            _claimed_only(results, print_found, print_all)
            and results[site]["status"].status != QueryStatus.CLAIMED
        ):
            continue
        query_time = results[site]["status"].query_time
        row = [
            username,
            site,
            results[site]["url_main"] or "",
            results[site]["url_user"] or "",
            str(results[site]["status"].status),
            results[site]["http_status"],
            "" if query_time is None else query_time,
        ]
        ws.append(row)
        r = ws.max_row
        for col, url in ((3, results[site]["url_main"]), (4, results[site]["url_user"])):
            if url:
                cell = ws.cell(row=r, column=col)
                cell.hyperlink = url
                cell.font = link_font
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    wb.save(path)


def write_jsonl(path: str, username: str, results: dict[str, dict[str, Any]]) -> None:
    with open(path, "w", encoding="utf-8") as handle:
        for site, data in results.items():
            status = data.get("status")
            record = {
                "username": username,
                "site": site,
                "url_main": data.get("url_main"),
                "url_user": data.get("url_user"),
                "status": str(status.status) if status is not None else None,
                "http_status": data.get("http_status"),
                "query_time": status.query_time if status is not None else None,
                "context": status.context if status is not None else None,
                "profile": data.get("profile"),
                "p_profile": data.get("p_profile", getattr(status, "p_profile", None)),
                "score_reasons": data.get("score_reasons", list(getattr(status, "score_reasons", ()) or ())),
            }
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")
